"""Download/normalize official BNS workbooks, preserving originals and provenance.

The selected sources are annual *internal* migration, total population, both sexes.
No urban/rural or sex subtotals are added to the already aggregated total.
"""
from pathlib import Path
from datetime import datetime,timezone
import argparse
import hashlib
import io
import json
import re
import sys

import openpyxl
import pandas as pd
import requests
from bs4 import BeautifulSoup

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from analytics import validate_panel

SOURCES={
 'arrivals':{'id':449811,'code':613301,'title':'Барлық ағындар бойынша ішкі көші-қон - келуі (жылдар бойынша)'},
 'departures':{'id':449812,'code':613302,'title':'Барлық ағындар бойынша ішкі көші-қон - кетуі (жылдар бойынша)'},
 'balance':{'id':449810,'code':613303,'title':'Барлық ағындар бойынша ішкі көші-қон айырымы (жылдар бойынша)'},
 'population':{'id':6644,'code':611104,'title':'Халықтың орташа жылдық саны (жылдар бойынша)'}
}
KATO={100000000:'abai',110000000:'akmola',150000000:'aktobe',190000000:'almaty_region',230000000:'atyrau',270000000:'west_kz',310000000:'zhambyl',330000000:'zhetisu',350000000:'karaganda',390000000:'kostanay',430000000:'kyzylorda',470000000:'mangystau',550000000:'pavlodar',590000000:'north_kz',610000000:'turkistan',620000000:'ulytau',630000000:'east_kz',710000000:'astana',750000000:'almaty',790000000:'shymkent'}


def parse_workbook(path,metric):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    metadata={str(row[0]).strip():row[1] for row in wb['Метадеректер'].iter_rows(values_only=True) if row[0] is not None}
    if int(metadata['Статистикалық көрсеткіштің коды'])!=SOURCES[metric]['code']:
        raise ValueError(f'Unexpected indicator code in {metric}; review source schema.')
    if '2009 жылдан' not in metadata.get('Ескертпе','') or 'қайта есептелді' not in metadata.get('Ескертпе',''):
        raise ValueError('No confirmed official territorial backcast; review comparability before importing.')
    sheet=next(s for s in wb if s.title.startswith('2009-')) if metric=='population' else wb['Көрсеткіш']
    rows=list(sheet.iter_rows(values_only=True))
    header=next(i for i,r in enumerate(rows) if str(r[0]).strip()=='ӘАОЖ коды')
    years={int(re.match(r'^(\d{4})',str(v)).group(1)):i for i,v in enumerate(rows[header]) if v is not None and re.match(r'^(19|20)\d{2}',str(v))}
    national={}
    result={}
    for row in rows[header+1:]:
        if row[1]=='Қазақстан Республикасы' and not national:
            national={year:row[index] for year,index in years.items()}
        try: code=int(row[0])
        except (TypeError,ValueError): continue
        if code not in KATO: continue
        region=KATO[code]
        if region in result: break # End total population block, not urban/sex subtotals.
        result[region]={year:row[index] for year,index in years.items()}
        if len(result)==20: break
    if len(result)!=20 or not national: raise ValueError('Expected 20 regions and national total.')
    return result,national,metadata,sheet.title


def build(refresh=False):
    raw=ROOT/'data/raw';raw.mkdir(exist_ok=True)
    fetched=datetime.now(timezone.utc).isoformat()
    page_url='https://stat.gov.kz/industries/social-statistics/demography/dynamic-tables/'
    page=requests.get(page_url,timeout=30);page.raise_for_status()
    soup=BeautifulSoup(page.text,'html.parser')
    parsed={}; national={}; provenance=[]; staged={}
    for metric,config in SOURCES.items():
        url=f'https://stat.gov.kz/api/iblock/element/{config["id"]}/file/kk/'
        path=raw/(metric+'.xlsx')
        if refresh or not path.exists():
            response=requests.get(url,timeout=45);response.raise_for_status();blob=response.content
        else:
            blob=path.read_bytes()
        staged[path]=blob
        values,totals,metadata,sheet=parse_workbook(io.BytesIO(blob),metric)
        parsed[metric]=values;national[metric]=totals
        row=soup.find(id=re.compile('_'+str(config['id'])+'$'))
        cells=row.select('.divTableCell') if row else []
        provenance.append({'metric':metric,'title':config['title'],'indicator_code':config['code'],'url':url,
            'updated_on':cells[1].get_text(strip=True) if len(cells)>1 else None,
            'file':'data/raw/'+path.name,'sheet':sheet,'sha256':hashlib.sha256(blob).hexdigest(),
            'definition':metadata.get('Көрсеткіштің анықтамасы'),'boundary_note':metadata.get('Ескертпе')})
    available=set.intersection(*(set(national[m]) for m in SOURCES))
    complete=[year for year in sorted(available) if year>=2009 and all(isinstance(values[year],(int,float)) for m in SOURCES for values in parsed[m].values())]
    if len(complete)<4: raise ValueError('Insufficient complete officially backcast years.')
    years=complete[-10:]
    rows=[]
    for year in years:
        for rid in KATO.values():
            row={'year':year,'region_id':rid,**{m:parsed[m][rid][year] for m in ['population','arrivals','departures']}}
            if row['arrivals']-row['departures']!=parsed['balance'][rid][year]:
                raise ValueError(f'Published balance discrepancy: {year}/{rid}')
            rows.append(row)
        for metric in SOURCES:
            total=sum(parsed[metric][rid][year] for rid in KATO.values())
            if abs(total-national[metric][year])>1:
                raise ValueError(f'National/region total mismatch: {metric}/{year}')
        if national['arrivals'][year]!=national['departures'][year]:
            raise ValueError(f'Internal migration national balance is not zero: {year}')
    frame=validate_panel(pd.DataFrame(rows))
    csv=frame.to_csv(index=False).encode()
    meta={'kind':'official','title':'Ұлттық статистика бюросының ресми деректері',
          'source':f'stat.gov.kz · {years[0]}–{years[-1]} · ішкі көші-қон, барлық халық',
          'url':page_url,'retrieved_at':fetched,'years':years,'sha256':hashlib.sha256(csv).hexdigest(),
          'note':'ҰСБ метадерегі бойынша 2009 жылдан бері қатарлар қазіргі облыстар мен Шымкент қаласының құрылуына сәйкес қайта есептелген. Соңғы 10 толық жыл алынды. Халық саны — жылдық орташа; .5 мәндері дөңгелектелмеген.',
          'source_caveat':'Халықтың орташа жылдық саны файлының кесте ішінде «жыл басына» және 611101 деген қайшы жазулар бар. Файл атауы мен Метадеректер парағының 611104 коды және анықтамасы орташа жылдық халықты көрсетеді; есептеуде осы метадерек басшылыққа алынды.',
          'checks':{'region_year_rows':len(frame),'national_totals_match':True,'published_balances_match':True,'national_internal_balance_zero':True},
          'sources':provenance}
    for path,blob in staged.items():
        path.write_bytes(blob)
    (ROOT/'data/official.csv').write_bytes(csv)
    (ROOT/'data/official-metadata.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'years':years,'rows':len(frame),'totals':frame.groupby('year')[['arrivals','departures','population']].sum().to_dict('index'),'checks':meta['checks']},ensure_ascii=False,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--refresh',action='store_true',help='Re-download official workbooks over verified HTTPS')
    build(parser.parse_args().refresh)
